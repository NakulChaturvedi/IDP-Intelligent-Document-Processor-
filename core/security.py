"""
core/security.py
API authentication and usage tracking.
"""

from __future__ import annotations

import hashlib
import json
import secrets
import threading
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import Depends, HTTPException, Request
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer

# ── Constants ─────────────────────────────────────────────────────────────────
# Resolve paths relative to THIS file (core/security.py → project root is parent)
_ROOT       = Path(__file__).parent.parent.resolve()
TOKENS_FILE = _ROOT / "tokens.json"
USAGE_LOG   = _ROOT / "usage_log.xlsx"

print(f"[security] Project root : {_ROOT}")
print(f"[security] Tokens file  : {TOKENS_FILE}")
print(f"[security] Usage log    : {USAGE_LOG}")

# ── Thread locks ──────────────────────────────────────────────────────────────
_excel_lock  = threading.Lock()
_tokens_lock = threading.Lock()

# ── Auth scheme ───────────────────────────────────────────────────────────────
_bearer = HTTPBearer(auto_error=False)


# ── Token store ───────────────────────────────────────────────────────────────

def _load_token_store() -> dict:
    if not TOKENS_FILE.exists():
        return {}
    try:
        return json.loads(TOKENS_FILE.read_text())
    except Exception:
        return {}


def _save_token_store(store: dict) -> None:
    TOKENS_FILE.write_text(json.dumps(store, indent=2))


def _hash(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


# ── Public: generate a token ──────────────────────────────────────────────────

def generate_token(label: str) -> str:
    token      = f"po-{secrets.token_urlsafe(32)}"
    token_hash = _hash(token)

    with _tokens_lock:
        store = _load_token_store()
        store[token_hash] = {
            "label":      label,
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        _save_token_store(store)

    print(f"[security] Token generated for '{label}' → {TOKENS_FILE}")
    return token

from cryptography.fernet import Fernet

def generate_encryption_key() -> str:
    """Run once to generate your encryption key. Store it safely."""
    return Fernet.generate_key().decode()

def encrypt_token(token: str, key: str) -> str:
    """Encrypt a token for safe storage in Excel log."""
    f = Fernet(key.encode())
    return f.encrypt(token.encode()).decode()

def decrypt_token(encrypted: str, key: str) -> str:
    """Decrypt a token from the Excel log."""
    f = Fernet(key.encode())
    return f.decrypt(encrypted.encode()).decode()

# ── Public: FastAPI dependency ────────────────────────────────────────────────

def verify_token(
    request:     Request,
    credentials: Optional[HTTPAuthorizationCredentials] = Depends(_bearer),
) -> dict:
    cfg = getattr(request.app.state, "cfg", {})

    if not cfg.get("security", {}).get("enabled", True):
        return {"token": "dev", "label": "anonymous"}

    if not credentials:
        raise HTTPException(
            status_code=401,
            detail="Missing authentication token. Include 'Authorization: Bearer <token>' header.",
        )

    plain = credentials.credentials
    h     = _hash(plain)
    store = _load_token_store()

    if h not in store:
        raise HTTPException(status_code=401, detail="Invalid authentication token.")

    return {"token": plain, "label": store[h]["label"]}


# ── Public: log one API call ──────────────────────────────────────────────────

def log_usage(
    *,
    token:      str,
    user:       str,
    endpoint:   str,
    model:      str,
    po_file:    str   = "",
    status:     str,
    duration:   float,
    fields:     int   = 0,
    line_items: int   = 0,
    error:      str   = "",
) -> None:
    row = {
        "Timestamp":      datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Token (masked)": _mask(token),
        "User":           user,
        "Endpoint":       endpoint,
        "Model":          model,
        "PO File":        Path(po_file).name if po_file else "",
        "Status":         status,
        "Duration (s)":   round(duration, 2),
        "Fields Filled":  fields,
        "Line Items":     line_items,
        "Error":          error[:300] if error else "",
    }

    print(f"[security] Logging usage → {USAGE_LOG}  (status={status})")

    with _excel_lock:
        _append_to_excel(USAGE_LOG, row)

def _mask(token: str) -> str:
    if len(token) <= 10:
        return token
    return token[:10] + "..."


# ── Internal: Excel append ────────────────────────────────────────────────────

_HEADERS = [
    "Timestamp", "Token (masked)", "User", "Endpoint", "Model",
    "PO File", "Status", "Duration (s)", "Fields Filled", "Line Items", "Error",
]

_COL_WIDTHS  = [20, 18, 15, 25, 28, 35, 10, 13, 14, 12, 45]
_HEADER_FILL = "1F3864"
_SUCCESS_BG  = "C6EFCE"
_SUCCESS_FG  = "276221"
_ERROR_BG    = "FFC7CE"
_ERROR_FG    = "9C0006"


def _append_to_excel(log_path: Path, row: dict) -> None:
    try:
        from openpyxl import load_workbook, Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        if log_path.exists():
            wb = load_workbook(log_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Usage Log"

            for col, header in enumerate(_HEADERS, 1):
                cell           = ws.cell(1, col, header)
                cell.font      = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill      = PatternFill("solid", start_color=_HEADER_FILL)
                cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[1].height = 18

            for col, width in enumerate(_COL_WIDTHS, 1):
                ws.column_dimensions[ws.cell(1, col).column_letter].width = width

            ws.freeze_panes = "A2"

        next_row = ws.max_row + 1
        for col, key in enumerate(_HEADERS, 1):
            cell      = ws.cell(next_row, col, row.get(key, ""))
            cell.font = Font(name="Arial", size=10)

        status_col  = _HEADERS.index("Status") + 1
        status_cell = ws.cell(next_row, status_col)
        if row["Status"] == "success":
            status_cell.fill = PatternFill("solid", start_color=_SUCCESS_BG)
            status_cell.font = Font(name="Arial", size=10, bold=True, color=_SUCCESS_FG)
        else:
            status_cell.fill = PatternFill("solid", start_color=_ERROR_BG)
            status_cell.font = Font(name="Arial", size=10, bold=True, color=_ERROR_FG)

        wb.save(log_path)
        print(f"[security] ✓ Usage log written ({next_row - 1} rows) → {log_path}")

    except Exception as e:
        print(f"[security] ✗ FAILED to write usage log: {e}")
        import traceback; traceback.print_exc()


# ── Public: usage summary ─────────────────────────────────────────────────────

def get_usage_summary() -> dict:
    if not USAGE_LOG.exists():
        return {"total_calls": 0, "by_user": {}}

    try:
        from openpyxl import load_workbook
        wb   = load_workbook(USAGE_LOG, read_only=True, data_only=True)
        ws   = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        by_user: dict[str, dict] = {}
        for row in rows:
            record = dict(zip(_HEADERS, row))
            user   = record.get("User") or "unknown"
            entry  = by_user.setdefault(user, {
                "calls": 0, "success": 0, "errors": 0,
                "total_fields": 0, "total_line_items": 0, "total_duration": 0.0,
            })
            entry["calls"]          += 1
            entry["total_duration"] += float(record.get("Duration (s)") or 0)
            if record.get("Status") == "success":
                entry["success"]          += 1
                entry["total_fields"]     += int(record.get("Fields Filled") or 0)
                entry["total_line_items"] += int(record.get("Line Items") or 0)
            else:
                entry["errors"] += 1

        return {"total_calls": len(rows), "by_user": by_user}

    except Exception as e:
        print(f"[security] WARNING: Failed to read usage log: {e}")
        return {"total_calls": 0, "by_user": {}}