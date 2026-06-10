"""
core/template.py
Auto-scans a customer Excel template to discover header fields and line-item columns.
Template matching uses company name extracted from the PO (Ship To field).
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


def normalize(text: str) -> str:
    return " ".join(text.split()) if text else ""


def label_to_key(label: str) -> str:
    key = label.lower()
    key = re.sub(r"[^a-z0-9]+", "_", key)
    return key.strip("_")


def is_skipped(label: str, skip_labels: set, skip_prefixes: tuple) -> bool:
    if label in skip_labels:
        return True
    if label.startswith(skip_prefixes):
        return True
    return False


def scan_template(template_path: str, cfg: dict) -> dict:
    """
    Scan a customer Excel template and return all field positions.
    Returns:
        header_fields:  [(key, label, row, value_col, conf_col)]
        li_fields:      [(key, label, col_number)]
        li_header_row:  int | None
    """
    from openpyxl import load_workbook

    sheet_name    = cfg["template"]["sheet_name"]
    anchor        = cfg["template"]["line_items_anchor"]
    skip_labels   = set(cfg["template"]["skip_labels"])
    skip_prefixes = tuple(cfg["template"]["skip_prefixes"])

    LABEL_TO_VALUE = {1: 2, 4: 5}

    wb = load_workbook(template_path, read_only=True)
    ws = wb[sheet_name]

    li_header_row = None
    header_fields = []
    li_fields     = []
    seen_keys     = set()

    for row in ws.iter_rows():
        for cell in row:
            if not cell.value or not isinstance(cell.value, str):
                continue

            label = normalize(cell.value)
            if not label or is_skipped(label, skip_labels, skip_prefixes):
                continue

            if label == normalize(anchor):
                li_header_row = cell.row

            if li_header_row and cell.row == li_header_row:
                key = label_to_key(label)
                li_fields.append((key, label, cell.column))

            if cell.column in LABEL_TO_VALUE:
                if li_header_row and cell.row >= li_header_row:
                    continue
                key = label_to_key(label)
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                val_col  = LABEL_TO_VALUE[cell.column]
                conf_col = val_col + 1
                header_fields.append((key, label, cell.row, val_col, conf_col))

    wb.close()
    li_fields.sort(key=lambda x: x[2])

    print(f"  Template scanned: {len(header_fields)} header field(s), {len(li_fields)} line-item column(s)")
    return {
        "header_fields": header_fields,
        "li_fields":     li_fields,
        "li_header_row": li_header_row,
    }


def find_template_by_company(company_name: str, template_dir: str) -> dict:
    """
    Fuzzy-match company name against template filenames in template_dir.
    Returns a dict with matched template info.
    """
    td       = Path(template_dir)
    all_xlsx = list(td.rglob("*.xlsx"))
    all_names = [f.name for f in all_xlsx]

    if not all_xlsx:
        return {"matched": False, "template_path": None, "template_name": None,
                "score": 0, "all_templates": []}

    stopwords = {"template", "the", "inc", "ltd", "llc", "co", "and",
                 "company", "associates", "&"}

    def words(text):
        parts = re.sub(r"[^a-z0-9]", " ", text.lower()).split()
        return [w for w in parts if w not in stopwords]

    def score(company, stem):
        cn = words(company)
        ts = words(stem)
        if not cn or not ts:
            return 0
        return sum(1 for w in cn if w in ts) + sum(1 for w in ts if w in cn)

    scored = sorted(all_xlsx, key=lambda f: score(company_name, f.stem), reverse=True)
    best   = scored[0]
    best_score = score(company_name, best.stem)

    return {
        "matched":       best_score > 0,
        "template_path": str(best) if best_score > 0 else None,
        "template_name": best.name if best_score > 0 else None,
        "score":         best_score,
        "all_templates": all_names,
    }


def find_template_by_po(po_filename: str, template_dir: str) -> str | None:
    """Legacy: match template by PO filename (used when company name not available)."""
    td      = Path(template_dir)
    po_stem = Path(po_filename).stem.lower()
    all_xlsx = list(td.rglob("*.xlsx"))

    for f in all_xlsx:
        if f.stem.lower() == po_stem:
            return str(f)
    for f in all_xlsx:
        if f.stem.lower() in po_stem or po_stem in f.stem.lower():
            return str(f)
    if len(all_xlsx) == 1:
        return str(all_xlsx[0])
    return None


def _loose(text: str) -> str:
    import re
    return re.sub(r"[^a-z0-9]", "", text.lower())
